#!/usr/bin/env python3
"""Export the current HSK Japanese glosses for a human to review in Excel."""

import argparse
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from hsk_workbook import HEADERS, REVIEWER_SHEET, current_rows


def text_cell(cell, value):
    cell.value = value
    # Treat every supplied string as literal text, including formula prefixes.
    cell.data_type = "s"


def export_workbook(output, root=None):
    rows = current_rows() if root is None else current_rows(root)
    book = Workbook()
    book.properties.creator = "Chabiko"
    book.properties.created = book.properties.modified = datetime(2000, 1, 1)
    instructions = book.active
    instructions.title = "記入方法"
    for line in [
        "HSK 1〜4 日本語訳の確認",
        "日本語訳はAIによる下書きです。日本語の意味・自然さをご確認ください。",
        "各HSKシートの判定欄で「OK」「修正」「要相談」を選んでください。",
        "OK：現在の日本語訳を承認します。修正後の日本語訳は空欄にしてください。",
        "修正：修正後の日本語訳に、承認する訳の全文を入力してください。",
        "要相談：承認せず、気になる点をコメントに記入してください。",
        "判定が空欄の行は未確認です。入力途中の訳やコメントだけでは承認されません。",
        "確認者シートの氏名・役割・確認日を、ご本人が入力してください。",
        "役割は human-language-reviewer と入力してください（日本語訳の確認担当）。",
        "確認日は YYYY-MM-DD 形式で入力してください。氏名等は確認記録として保存されます。",
        "確認範囲は日本語訳のみです。字体・発音・権利の承認にはなりません。",
        "ID・簡体字・ピンイン・現在の日本語訳は変更しないでください。ID列は非表示です。",
        "行やシートを削除しないでください。フィルターで絞り込んで作業できます。",
        "元データの更新があった場合は取込できません。担当者から新しいファイルを受け取ってください。",
        "記入したファイルを .xlsx 形式で保存し、担当者に返送してください。",
        "出典：簡体字・ピンイン・HSKレベルは HearMandarin（https://hearmandarin.com/datasets/）の HSK 3.0 単語リストを使用しています。",
        "ライセンス：CC BY 4.0（https://creativecommons.org/licenses/by/4.0/）。元データから一部の項目のみを使用し、日本語訳は独自に追加したものです。HearMandarin は HSK 試験の主催者とは無関係です。",
    ]:
        instructions.append([line])
    instructions.column_dimensions["A"].width = 105
    for row in instructions:
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        instructions.row_dimensions[row[0].row].height = 34
    reviewer = book.create_sheet(REVIEWER_SHEET)
    for label in ["氏名", "役割", "確認日"]:
        reviewer.append([label, None])
        reviewer.cell(reviewer.max_row, 2).number_format = "@"
    reviewer.column_dimensions["A"].width = 14
    reviewer.column_dimensions["B"].width = 45
    for level in range(1, 5):
        sheet = book.create_sheet(f"HSK{level}")
        sheet.append(HEADERS)
        for record in rows.values():
            if record["hsk"]["introducedAtLevel"] == level:
                sheet.append([None] * 7)
                for index, value in enumerate([record["id"], record["simplified"], record["pinyin"], record["japanese"]], 1):
                    text_cell(sheet.cell(sheet.max_row, index), value)
        sheet.freeze_panes = "E2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions["A"].hidden = True
        for column, width in zip("ABCDEFG", [25, 18, 24, 44, 12, 44, 44]):
            sheet.column_dimensions[column].width = width
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.protection = Protection(locked=cell.row == 1 or cell.column <= 4)
                if cell.row == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="275D55")
            sheet.row_dimensions[row[0].row].height = 36
        # Protection is an editing aid, not an integrity/security boundary.
        sheet.protection.sheet = True
        sheet.protection.autoFilter = False
        sheet.protection.selectLockedCells = True
        validation = DataValidation(type="list", formula1='"OK,修正,要相談"', allow_blank=True)
        validation.errorTitle = "判定を選んでください"
        validation.error = "OK・修正・要相談のいずれかを選んでください。"
        validation.showErrorMessage = True
        validation.errorStyle = "stop"
        sheet.add_data_validation(validation)
        validation.add(f"E2:E{sheet.max_row}")
    buffer = BytesIO()
    book.save(buffer)
    # openpyxl uses wall-clock ZIP timestamps and resets modified on save.
    with ZipFile(buffer) as source, ZipFile(output, "w", compression=ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            data = source.read(name)
            if name == "docProps/core.xml":
                import re
                data = re.sub(rb"<dcterms:modified[^>]*>.*?</dcterms:modified>",
                              b'<dcterms:modified xmlns:dcterms="http://purl.org/dc/terms/" xsi:type="dcterms:W3CDTF">2000-01-01T00:00:00Z</dcterms:modified>', data)
            entry = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            target.writestr(entry, data)
    book.close()
    return len(rows)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(f"Exported {export_workbook(args.output)} rows to {args.output}")
