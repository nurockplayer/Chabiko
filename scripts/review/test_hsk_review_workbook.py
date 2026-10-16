"""Exercise the real workbook CLIs against temporary repository copies."""

from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook

from hsk_workbook import ROOT, current_rows, importer


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="chabiko-workbook-test-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        for relative in (importer.PROFILE_REL.parent, Path("data/hsk-vocabulary"), Path("scripts/review")):
            shutil.copytree(ROOT / relative, self.root / relative, ignore=shutil.ignore_patterns("__pycache__"))
        shutil.copy2(ROOT / "scripts/import-hearmandarin-hsk-json.py", self.root / "scripts")
        self.workbook = self.root / "review.xlsx"
        self.cli("export", "--output", self.workbook)

    def cli(self, kind, *args, expected=0):
        result = subprocess.run([sys.executable, str(self.root / f"scripts/review/{kind}_hsk_review_workbook.py"), *map(str, args)], capture_output=True, text=True)
        self.assertEqual(result.returncode, expected, result.stdout + result.stderr)
        return result.stdout + result.stderr

    def fill(self, change=None, reviewer=True):
        book = load_workbook(self.workbook)
        if reviewer:
            for index, item in enumerate([" Test reviewer (synthetic fixture) ", "human-language-reviewer", "2026-10-16"], 1):
                book["確認者"].cell(index, 2, item)
        if change:
            change(book)
        book.save(self.workbook)
        book.close()

    def snapshot(self):
        return {str(path.relative_to(self.root)): path.read_bytes() for path in (self.root / "data").rglob("*") if path.is_file()}

    def test_deterministic_export_and_format(self):
        second = self.root / "second.xlsx"
        self.cli("export", "--output", second)
        self.assertEqual(self.workbook.read_bytes(), second.read_bytes())
        book = load_workbook(self.workbook)
        self.assertEqual(sum(book[f"HSK{i}"].max_row - 1 for i in range(1, 5)), 1969)
        self.assertEqual([book["確認者"].cell(i, 2).value for i in range(1, 4)], [None] * 3)
        notes = "\n".join(str(row[0].value) for row in book["記入方法"].iter_rows())
        self.assertIn("HearMandarin", notes)
        self.assertIn("CC BY 4.0", notes)
        for level in range(1, 5):
            sheet = book[f"HSK{level}"]
            self.assertTrue(sheet.column_dimensions["A"].hidden)
            self.assertTrue(sheet.protection.sheet)
            self.assertFalse(sheet["E2"].protection.locked)
            self.assertEqual(sheet.freeze_panes, "E2")
            self.assertEqual(next(iter(sheet.data_validations.dataValidation)).formula1, '"OK,修正,要相談"')
        book.close()

    def test_round_trip_dry_run_and_rebuild(self):
        def change(book):
            sheet = book["HSK1"]
            sheet["E2"], sheet["E3"], sheet["F3"] = "OK", "修正", "八（はち）。"
            sheet["E4"], sheet["G4"] = "要相談", "テスト用の相談"
            sheet["G5"] = "判定なしのコメント"
            book["HSK2"]["E2"] = "OK"
        self.fill(change)
        before = self.snapshot()
        self.assertIn("Dry run", self.cli("import", "--input", self.workbook, "--dry-run"))
        self.assertEqual(before, self.snapshot())
        self.cli("import", "--input", self.workbook)
        rows = list(current_rows(self.root).values())
        self.assertEqual([row["reviewStatus"] for row in rows[:4]], ["reviewed", "reviewed", "draft", "draft"])
        self.assertEqual(rows[1]["japanese"], "八（はち）。")
        self.assertIn("Japanese glosses are human reviewed.", rows[1]["source"]["note"])
        self.assertIn("not human reviewed", rows[2]["source"]["note"])
        ledger = importer.load_json(self.root / importer.HUMAN_REVIEWS_REL)
        artifact = ledger["artifacts"][0]
        self.assertEqual(artifact["reviewerIdentity"], " Test reviewer (synthetic fixture) ")
        self.assertEqual(artifact["outcomeCounts"], {"OK": 2, "修正": 1, "要相談": 1, "blank": 1965})
        output = self.root / "output"
        subprocess.run([sys.executable, str(self.root / "scripts/import-hearmandarin-hsk-json.py"), "--from-manifest", "--output-dir", str(output)], check=True, capture_output=True)
        manifest = importer.load_json(output / "manifest.json")
        self.assertEqual((output / "manifest.json").read_bytes(), (self.root / importer.PROFILE_REL.parent / "manifest.json").read_bytes())
        for batch in manifest["publication"]["batches"]:
            self.assertEqual((output / batch["file"]).read_bytes(), (self.root / "data/hsk-vocabulary" / batch["file"]).read_bytes())
        # A new partial review preserves earlier approvals on blank rows and
        # explicitly returns a previously accepted row to draft for consultation.
        self.cli("export", "--output", self.workbook)
        self.fill(lambda book: setattr(book["HSK1"]["E3"], "value", "要相談"))
        self.cli("import", "--input", self.workbook)
        later_rows = list(current_rows(self.root).values())
        self.assertEqual([row["reviewStatus"] for row in later_rows[:2]], ["reviewed", "draft"])
        # A later edit cannot inherit the old human approval.
        companion_path = self.root / importer.JAPANESE_REL
        companion = importer.load_json(companion_path)
        companion["records"][0]["japanese"] = "Changed after review"
        companion_path.write_bytes(importer.json_bytes(companion))
        output2 = self.root / "output2"
        subprocess.run([sys.executable, str(self.root / "scripts/import-hearmandarin-hsk-json.py"), "--from-manifest", "--output-dir", str(output2)], check=True, capture_output=True)
        self.assertEqual(importer.load_json(output2 / importer.FIRST_BATCH_FILENAME)["vocabulary"][0]["reviewStatus"], "draft")

    def test_refusals_are_no_ops(self):
        original = self.workbook.read_bytes()
        cases = [
            (False, lambda book: None, "reviewer name"),
            (True, lambda book: setattr(book["確認者"]["B2"], "value", "maintainer"), "reviewer role"),
            (True, lambda book: setattr(book["確認者"]["B3"], "value", "2026-02-30"), "error:"),
            (True, lambda book: setattr(book["HSK1"]["D2"], "value", "stale"), "HSK1 row 2: stale"),
            (True, lambda book: setattr(book["HSK1"]["A2"], "value", "unknown"), "unknown ID"),
            (True, lambda book: setattr(book["HSK1"]["A3"], "value", book["HSK1"]["A2"].value), "duplicate ID"),
            (True, lambda book: book["HSK1"].delete_rows(2), "missing 1 IDs"),
            (True, lambda book: setattr(book["HSK1"]["E2"], "value", "修正"), "nonempty corrected gloss"),
            (True, lambda book: setattr(book["HSK1"]["E2"], "value", "maybe"), "invalid decision"),
            (True, lambda book: setattr(book["HSK1"]["F2"], "value", "=1+1"), "formula/error"),
        ]
        before = self.snapshot()
        for reviewer, change, expected in cases:
            with self.subTest(expected=expected):
                self.workbook.write_bytes(original)
                self.fill(change, reviewer)
                self.assertIn(expected, self.cli("import", "--input", self.workbook, expected=2))
                self.assertEqual(before, self.snapshot())

    def test_unrebuilt_companion_edit_is_not_overwritten(self):
        self.fill(lambda book: setattr(book["HSK1"]["E2"], "value", "OK"))
        path = self.root / importer.JAPANESE_REL
        companion = importer.load_json(path)
        companion["records"][0]["japanese"] = "New author edit"
        path.write_bytes(importer.json_bytes(companion))
        before = self.snapshot()
        self.assertIn("out of sync", self.cli("import", "--input", self.workbook, expected=2))
        self.assertEqual(before, self.snapshot())


if __name__ == "__main__":
    unittest.main()
