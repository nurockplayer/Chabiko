#!/usr/bin/env python3
"""Validate a returned human HSK workbook, then rebuild companions and batches."""

import argparse
from contextlib import redirect_stdout
from copy import deepcopy
from datetime import date, datetime
from io import StringIO
from pathlib import Path
import shutil
import sys
import tempfile

from openpyxl import load_workbook

from hsk_workbook import HEADERS, REVIEWER_SHEET, ROOT, current_rows, importer


def value(cell):
    importer.require(cell.data_type not in ("f", "e"), f"formula/error cell is not allowed: {cell.coordinate}")
    return cell.value


def optional_text(cell):
    result = value(cell)
    importer.require(result is None or isinstance(result, str), f"expected text at {cell.coordinate}")
    return result or ""


def read_workbook(path, rows):
    book = load_workbook(path, data_only=False)
    try:
        importer.require(book.sheetnames == ["記入方法", REVIEWER_SHEET, "HSK1", "HSK2", "HSK3", "HSK4"], "unexpected workbook sheets")
        sheet = book[REVIEWER_SHEET]
        importer.require([sheet.cell(i, 1).value for i in range(1, 4)] == ["氏名", "役割", "確認日"], "invalid reviewer fields")
        name, role, reviewed_at = [value(sheet.cell(i, 2)) for i in range(1, 4)]
        importer.require(isinstance(name, str) and bool(name.strip()), "reviewer name (氏名) is required")
        importer.require(role == "human-language-reviewer", "reviewer role (役割) must be human-language-reviewer")
        if isinstance(reviewed_at, datetime):
            importer.require(reviewed_at.time().isoformat() == "00:00:00", "review date must not include a time")
            reviewed_at = reviewed_at.date().isoformat()
        elif isinstance(reviewed_at, date):
            reviewed_at = reviewed_at.isoformat()
        importer.require(isinstance(reviewed_at, str) and date.fromisoformat(reviewed_at).isoformat() == reviewed_at,
                         "review date (確認日) must be YYYY-MM-DD")
        artifact = {"reviewerIdentity": name, "reviewerRole": role, "reviewDate": reviewed_at,
                    "approvalScope": ["japanese-gloss", "review-status"],
                    "reviewVersion": importer.sha256(path.read_bytes()),
                    "reviewVersionKind": "returned-workbook-sha256",
                    "reviewOutcome": "accepted", "unresolvedIssues": [], "blockedContent": [],
                    "outcomeCounts": {"OK": 0, "修正": 0, "要相談": 0, "blank": 0}, "rows": []}
        seen, errors = set(), []
        for level in range(1, 5):
            sheet = book[f"HSK{level}"]
            importer.require([cell.value for cell in sheet[1]] == HEADERS, f"HSK{level}: unexpected headers")
            for cells in sheet.iter_rows(min_row=2, max_col=7):
                if all(cell.value is None for cell in cells):
                    continue
                location = f"HSK{level} row {cells[0].row}"
                try:
                    record_id, simplified, pinyin, japanese = [value(cell) for cell in cells[:4]]
                    importer.require(isinstance(record_id, str) and record_id in rows, f"unknown ID {record_id!r}")
                    importer.require(record_id not in seen, f"duplicate ID {record_id}")
                    seen.add(record_id)
                    before = rows[record_id]
                    importer.require(before["hsk"]["introducedAtLevel"] == level, f"wrong level for {record_id}")
                    importer.require([simplified, pinyin, japanese] == [before[key] for key in ("simplified", "pinyin", "japanese")],
                                     f"stale or edited source text for {record_id}")
                    decision, correction, comment = [optional_text(cell) for cell in cells[4:]]
                    importer.require(decision in ("", "OK", "修正", "要相談"), f"invalid decision {decision!r}")
                    importer.require(decision != "修正" or bool(correction.strip()), "修正 requires a nonempty corrected gloss")
                    importer.require(decision != "OK" or not correction, "OK must not have a corrected gloss; choose 修正")
                    artifact["outcomeCounts"][decision or "blank"] += 1
                    if decision or correction or comment:
                        before = {key: deepcopy(before[key]) for key in ("id", "simplified", "pinyin", "hsk", "japanese")}
                        after = deepcopy(before)
                        if decision == "修正":
                            after["japanese"] = correction
                        item = {"decision": decision, "before": before, "after": after,
                                "reviewVersion": importer.review_fingerprint(after),
                                "correction": correction, "comment": comment}
                        artifact["rows"].append(item)
                    if decision == "要相談":
                        artifact["unresolvedIssues"].append({"id": record_id, "comment": comment})
                        artifact["blockedContent"].append(record_id)
                except (importer.ImportFailure, TypeError) as error:
                    errors.append(f"{location}: {error}")
        missing = set(rows) - seen
        if missing:
            errors.append(f"missing {len(missing)} IDs, including {sorted(missing)[:5]}")
        importer.require(not errors, "\n".join(errors))
        if artifact["outcomeCounts"]["要相談"]:
            artifact["reviewOutcome"] = "needs-changes"
        elif not (artifact["outcomeCounts"]["OK"] + artifact["outcomeCounts"]["修正"]):
            artifact["reviewOutcome"] = "not-reviewed"
        return artifact
    finally:
        book.close()


def import_workbook(path, dry_run=False, root=ROOT):
    rows = current_rows(root)
    artifact = read_workbook(path, rows)
    # Blank decisions with notes remain in the receipt but never confer approval.
    ledger_path = root / importer.HUMAN_REVIEWS_REL
    ledger = importer.load_json(ledger_path) if ledger_path.exists() else {"schemaVersion": 1, "artifacts": []}
    if artifact in ledger["artifacts"]:
        print("Workbook already recorded; no changes.")
        return
    ledger["artifacts"].append(artifact)
    with tempfile.TemporaryDirectory(prefix="chabiko-hsk-review-") as temporary:
        stage = Path(temporary)
        package = importer.PROFILE_REL.parent
        shutil.copytree(root / package, stage / package)
        # A workbook must not overwrite a companion edit that has not yet been
        # rebuilt into the learner batches it was exported from.
        baseline = stage / "baseline"
        with redirect_stdout(StringIO()):
            importer.run_import(stage, None, baseline)
        importer.require((baseline / "manifest.json").read_bytes() == (root / package / "manifest.json").read_bytes(),
                         "repository manifest is out of sync; rebuild before importing")
        manifest = importer.load_json(stage / package / "manifest.json")
        for batch in manifest["publication"]["batches"]:
            importer.require((baseline / batch["file"]).read_bytes() == (root / "data/hsk-vocabulary" / batch["file"]).read_bytes(),
                             f"repository batch {batch['file']} is out of sync; rebuild before importing")
        (stage / importer.HUMAN_REVIEWS_REL).write_bytes(importer.json_bytes(ledger))
        corrections = {row["after"]["id"]: row["after"]["japanese"] for row in artifact["rows"] if row["decision"] == "修正"}
        for batch in manifest["publication"]["batches"]:
            placement = batch["placement"]
            companion_path = stage / importer._companion_path(placement)
            companion = importer.load_json(companion_path)
            changed = False
            for record in companion["records"]:
                record_id = f"hm-hsk3-{record['sourceId']}"
                if record_id in corrections:
                    record["japanese"] = corrections[record_id]
                    changed = True
            if changed:
                companion_path.write_bytes(importer.json_bytes(companion))
        output = stage / "rebuilt"
        with redirect_stdout(StringIO()):
            importer.run_import(stage, None, output)
        planned = {package / "manifest.json": (output / "manifest.json").read_bytes(),
                   importer.HUMAN_REVIEWS_REL: importer.json_bytes(ledger)}
        for batch in manifest["publication"]["batches"]:
            planned[Path("data/hsk-vocabulary") / batch["file"]] = (output / batch["file"]).read_bytes()
            companion = importer._companion_path(batch["placement"])
            planned[companion] = (stage / companion).read_bytes()
        changes = {root / rel: data for rel, data in planned.items() if not (root / rel).exists() or (root / rel).read_bytes() != data}
        print(f"{'Dry run' if dry_run else 'Import'}: {artifact['outcomeCounts']}; {len(changes)} files {'would change' if dry_run else 'changed'}.")
        if not dry_run:
            originals = {target: target.read_bytes() if target.exists() else None for target in changes}
            try:
                for target, data in changes.items():
                    target.write_bytes(data)
            except OSError:
                for target, original in originals.items():
                    if original is None:
                        target.unlink(missing_ok=True)
                    else:
                        target.write_bytes(original)
                raise


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    try:
        import_workbook(args.input, args.dry_run)
    except (importer.ImportFailure, OSError, KeyError, TypeError, ValueError) as error:
        print(f"error: {error}", file=sys.stderr)
        sys.exit(2)
